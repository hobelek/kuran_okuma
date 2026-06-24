import os
import openpyxl
from django.core.management.base import BaseCommand
from django.contrib.auth.models import User
from takip.models import Student

class Command(BaseCommand):
    help = 'Öğrencileri Excel dosyasından içe aktarır. Her sayfa bir sınıftır.'

    def add_arguments(self, parser):
        parser.add_argument('excel_path', type=str, help='Excel dosyasının yolu')

    def handle(self, *args, **options):
        excel_path = options['excel_path']
        if not os.path.exists(excel_path):
            self.stderr.write(self.style.ERROR(f'Dosya bulunamadı: {excel_path}'))
            return

        wb = openpyxl.load_workbook(excel_path)
        toplam_ogrenci = 0

        for sheet_name in wb.sheetnames:
            self.stdout.write(f'İşleniyor: {sheet_name} sayfası')
            sheet = wb[sheet_name]
            # Başlık satırını atla (No, İsim, Soyad)
            for row in sheet.iter_rows(min_row=2, values_only=True):
                if not row[0] or not row[1] or not row[2]:
                    continue
                student_no = str(row[0]).strip()
                first_name = str(row[1]).strip()
                last_name = str(row[2]).strip()
                classroom = sheet_name.strip()

                username = f'ogrenci_{student_no}'
                user, created = User.objects.get_or_create(
                    username=username,
                    defaults={
                        'first_name': first_name,
                        'last_name': last_name,
                        'is_superuser': False,
                        'is_staff': False
                    }
                )
                if not created:
                    user.first_name = first_name
                    user.last_name = last_name
                    user.save()

                student, std_created = Student.objects.get_or_create(
                    user=user,
                    defaults={
                        'student_no': student_no,
                        'first_name': first_name,
                        'last_name': last_name,
                        'classroom': classroom
                    }
                )
                if not std_created:
                    student.first_name = first_name
                    student.last_name = last_name
                    student.classroom = classroom
                    student.save()

                toplam_ogrenci += 1
                self.stdout.write(f'  - {student_no} {first_name} {last_name} ({classroom}) {"eklendi" if std_created else "güncellendi"}')

        self.stdout.write(self.style.SUCCESS(f'Toplam {toplam_ogrenci} öğrenci işlendi.'))